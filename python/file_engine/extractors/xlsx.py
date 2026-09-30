"""XLSX extraction via openpyxl, sheet by sheet."""
from __future__ import annotations

from pathlib import Path

from .base import ExtractionResult, register

MAX_ROWS_PER_SHEET = 5000


class XlsxExtractor:
    name = "xlsx"

    def supports(self, path: Path, extension: str) -> bool:
        return extension == ".xlsx"

    def extract(self, path: Path) -> ExtractionResult:
        result = ExtractionResult(extractor=self.name)
        try:
            from openpyxl import load_workbook
        except ImportError:
            result.error = "openpyxl is not installed"
            return result
        try:
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in workbook.worksheets:
                parts.append(f"--- Sheet: {sheet.title} ---")
                rows_written = 0
                for row in sheet.iter_rows(values_only=True):
                    parts.append(" | ".join("" if cell is None else str(cell) for cell in row))
                    rows_written += 1
                    if rows_written >= MAX_ROWS_PER_SHEET:
                        result.warnings.append(
                            f"Sheet {sheet.title!r} truncated at {MAX_ROWS_PER_SHEET} rows"
                        )
                        break
            result.sheetCount = len(workbook.sheetnames)
            workbook.close()
            result.text = "\n".join(parts)
            result.success = True
            result.wordCount = len(result.text.split())
            result.lineCount = result.text.count("\n") + 1
        except Exception as exc:  # noqa: BLE001
            result.error = f"XLSX extraction failed: {exc}"
        return result


register(XlsxExtractor())
