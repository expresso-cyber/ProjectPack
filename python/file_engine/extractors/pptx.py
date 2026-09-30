"""PPTX extraction via python-pptx, slide by slide."""
from __future__ import annotations

from pathlib import Path

from .base import ExtractionResult, register


class PptxExtractor:
    name = "pptx"

    def supports(self, path: Path, extension: str) -> bool:
        return extension == ".pptx"

    def extract(self, path: Path) -> ExtractionResult:
        result = ExtractionResult(extractor=self.name)
        try:
            from pptx import Presentation
        except ImportError:
            result.error = "python-pptx is not installed"
            return result
        try:
            presentation = Presentation(str(path))
            parts: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                parts.append(f"--- Slide {index} ---")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            line = "".join(run.text for run in para.runs)
                            if line:
                                parts.append(line)
                if getattr(slide, "has_notes_slide", False):
                    notes = slide.notes_slide.notes_text_frame.text
                    if notes:
                        parts.append(f"[Notes] {notes}")
            result.slideCount = len(presentation.slides)
            result.text = "\n".join(parts)
            result.success = True
            result.wordCount = len(result.text.split())
            result.lineCount = result.text.count("\n") + 1
        except Exception as exc:  # noqa: BLE001
            result.error = f"PPTX extraction failed: {exc}"
        return result


register(PptxExtractor())
