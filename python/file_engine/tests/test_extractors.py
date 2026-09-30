from file_engine.core.pipeline import extract_one


def test_extract_python(sample_project):
    res = extract_one(str(sample_project), "session01/s01.py", ".py")
    assert res["success"] is True
    assert "alpha hello world" in res["text"]
    assert res["extractor"] == "text"


def test_extract_csv(sample_project):
    res = extract_one(str(sample_project), "session01/data.csv", ".csv")
    assert res["success"] is True
    assert "apple" in res["text"]
    assert res["extractor"] == "csv"


def test_extract_binary_fails_cleanly(sample_project):
    res = extract_one(str(sample_project), "photo.bin", ".bin")
    assert res["success"] is False
    assert res["error"]


def test_extract_pdf(sample_project):
    import pymupdf
    pdf_path = sample_project / "doc.pdf"
    with pymupdf.open() as doc:
        doc.new_page()
        doc[0].insert_text((72, 72), "unique pdf words here")
        doc.save(str(pdf_path))
    res = extract_one(str(sample_project), "doc.pdf", ".pdf")
    assert res["success"] is True
    assert res["pageCount"] == 1
    assert "unique pdf words" in res["text"]


def test_extract_docx(sample_project):
    from docx import Document
    docx_path = sample_project / "notes.docx"
    doc = Document()
    doc.add_paragraph("docx paragraph one")
    doc.add_paragraph("docx paragraph two")
    doc.save(str(docx_path))
    res = extract_one(str(sample_project), "notes.docx", ".docx")
    assert res["success"] is True
    assert "docx paragraph one" in res["text"]


def test_extract_xlsx(sample_project):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["colA", "colB"])
    ws.append([1, 2])
    wb.save(str(sample_project / "sheet.xlsx"))
    res = extract_one(str(sample_project), "sheet.xlsx", ".xlsx")
    assert res["success"] is True
    assert res["sheetCount"] == 1
    assert "colA" in res["text"]


def test_extract_pptx(sample_project):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Slide One Title"
    prs.save(str(sample_project / "deck.pptx"))
    res = extract_one(str(sample_project), "deck.pptx", ".pptx")
    assert res["success"] is True
    assert res["slideCount"] == 1
    assert "Slide One Title" in res["text"]
