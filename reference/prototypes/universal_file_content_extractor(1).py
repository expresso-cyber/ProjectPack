from __future__ import annotations

import csv
import sys
from pathlib import Path
from typing import Optional

# ============================================================
# Configuration
# ============================================================

APP_NAME = "UNIVERSAL FILE CONTENT EXTRACTOR"
SEPARATOR = "=" * 70

# Extension -> display name
SUPPORTED_EXTENSIONS = {
    ".py": "Python",
    ".txt": "Text",
    ".md": "Markdown",
    ".html": "HTML",
    ".htm": "HTML",
    ".css": "CSS",
    ".js": "JavaScript",
    ".jsx": "JavaScript JSX",
    ".ts": "TypeScript",
    ".tsx": "TypeScript JSX",
    ".json": "JSON",
    ".xml": "XML",
    ".yaml": "YAML",
    ".yml": "YAML",
    ".sql": "SQL",
    ".csv": "CSV",
    ".log": "Log",
    ".ini": "INI",
    ".cfg": "Configuration",
    ".conf": "Configuration",
    ".java": "Java",
    ".c": "C",
    ".h": "C/C++ Header",
    ".cpp": "C++",
    ".hpp": "C++ Header",
    ".cs": "C#",
    ".php": "PHP",
    ".rb": "Ruby",
    ".go": "Go",
    ".rs": "Rust",
    ".sh": "Shell Script",
    ".bat": "Batch",
    ".ps1": "PowerShell",
    ".pdf": "PDF",
    ".docx": "Word Document",
    ".xlsx": "Excel Workbook",
    ".pptx": "PowerPoint Presentation",
}

TEXT_EXTENSIONS = {
    extension
    for extension, name in SUPPORTED_EXTENSIONS.items()
    if extension not in {".pdf", ".docx", ".xlsx", ".pptx"}
}

BINARY_EXTRACTORS = {".pdf", ".docx", ".xlsx", ".pptx"}


# ============================================================
# Console helpers
# ============================================================


def print_header() -> None:
    print("\n" + SEPARATOR)
    print(f"{APP_NAME:^70}")
    print(SEPARATOR + "\n")


def info(message: str) -> None:
    print(f"[INFO] {message}")


def success(message: str) -> None:
    print(f"[SUCCESS] {message}")


def warning(message: str) -> None:
    print(f"[WARNING] {message}")


def error(message: str) -> None:
    print(f"[ERROR] {message}")


def pause() -> None:
    input("\nPress Enter to continue...")


# ============================================================
# Input validation
# ============================================================


def clean_user_path(value: str) -> Path:
    """
    Remove common surrounding quotes from pasted Windows paths.
    """
    value = value.strip()

    if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
        value = value[1:-1].strip()

    return Path(value).expanduser()


def get_valid_source_folder() -> Path:
    while True:
        print("Enter the folder path:")
        print(r"Example: C:\Users\YourName\Documents\MyProject")
        raw_path = input("> ")

        if not raw_path.strip():
            error("Folder path cannot be empty.\n")
            continue

        path = clean_user_path(raw_path)

        try:
            path = path.resolve()
        except OSError:
            error("The path could not be resolved. Please try again.\n")
            continue

        if not path.exists():
            error("The specified folder does not exist.")
            print("Please enter a valid folder path.\n")
            continue

        if not path.is_dir():
            error("The specified path is not a folder.")
            print("Please enter a directory path instead.\n")
            continue

        return path


def get_valid_choice(prompt: str, minimum: int, maximum: int) -> int:
    while True:
        raw = input(prompt).strip()

        try:
            choice = int(raw)
        except ValueError:
            error(f"Please enter a number from {minimum} to {maximum}.")
            continue

        if minimum <= choice <= maximum:
            return choice

        error(f"Please enter a number from {minimum} to {maximum}.")


# ============================================================
# File discovery
# ============================================================


def scan_files(source_folder: Path) -> list[Path]:
    """
    Recursively scan all files. Permission/OS errors on individual
    entries are ignored so one inaccessible location does not stop
    the whole scan.
    """
    discovered: list[Path] = []

    try:
        for path in source_folder.rglob("*"):
            try:
                if path.is_file():
                    discovered.append(path)
            except OSError:
                continue
    except OSError as exc:
        error(f"Unable to complete folder scan: {exc}")

    return sorted(discovered, key=lambda p: str(p).lower())


def get_extension_groups(
    files: list[Path],
) -> dict[str, list[Path]]:
    groups: dict[str, list[Path]] = {}

    for file_path in files:
        extension = file_path.suffix.lower()

        if extension in SUPPORTED_EXTENSIONS:
            groups.setdefault(extension, []).append(file_path)

    return dict(sorted(groups.items()))


# ============================================================
# Menus
# ============================================================


def display_available_types(
    extension_groups: dict[str, list[Path]],
) -> list[str]:
    extensions = list(extension_groups.keys())

    print("\nSupported file types found in the selected folder:\n")

    for index, extension in enumerate(extensions, start=1):
        name = SUPPORTED_EXTENSIONS[extension]
        count = len(extension_groups[extension])
        plural = "file" if count == 1 else "files"
        print(f"{index:>2}. {name:<24} ({extension:<6}) → {count} {plural}")

    print(
        f"{len(extensions) + 1:>2}. "
        f"{'All supported file types':<24} → "
        f"{sum(len(paths) for paths in extension_groups.values())} files"
    )

    return extensions


def select_file_type(
    extension_groups: dict[str, list[Path]],
) -> Optional[list[Path]]:
    extensions = display_available_types(extension_groups)

    print()

    choice = get_valid_choice(
        "Select the file type you want to extract: ",
        1,
        len(extensions) + 1,
    )

    if choice == len(extensions) + 1:
        selected = [
            file_path
            for extension in extensions
            for file_path in extension_groups[extension]
        ]
        return selected

    selected_extension = extensions[choice - 1]
    return extension_groups[selected_extension]


def display_matching_files(
    source_folder: Path,
    files: list[Path],
    selection_title: str,
) -> None:
    print(f"\nFiles found for {selection_title}:\n")

    for index, file_path in enumerate(files, start=1):
        relative = file_path.relative_to(source_folder)
        print(f"{index:>4}. {relative}")

    count = len(files)
    print(f"\nTotal files: {count}")


def select_files(
    source_folder: Path,
    files: list[Path],
) -> Optional[list[Path]]:
    if not files:
        return None

    extensions = sorted({path.suffix.lower() for path in files})

    if len(extensions) == 1:
        title = extensions[0]
    else:
        title = "selected supported files"

    display_matching_files(source_folder, files, title)

    print("\nWhat would you like to do?\n")
    print(f"1. Extract content from all {len(files)} files")
    print("2. Select specific files")
    print("3. Cancel")

    choice = get_valid_choice("\nChoose an option: ", 1, 3)

    if choice == 1:
        return files

    if choice == 3:
        return None

    while True:
        raw = input(
            "\nEnter file numbers separated by commas " "(Example: 1,3,4): "
        ).strip()

        if not raw:
            error("Please enter at least one file number.")
            continue

        try:
            numbers = [int(part.strip()) for part in raw.split(",")]
        except ValueError:
            error("Invalid input. Use numbers separated by commas.")
            continue

        if len(set(numbers)) != len(numbers):
            error("Duplicate file numbers were entered.")
            continue

        if any(number < 1 or number > len(files) for number in numbers):
            error(f"Each number must be between 1 and {len(files)}.")
            continue

        return [files[number - 1] for number in numbers]


# ============================================================
# Extraction handlers
# ============================================================


def read_text_file(file_path: Path) -> str:
    """
    Read text robustly. UTF-8 is preferred; common fallback encodings
    are attempted before replacing undecodable characters.
    """
    encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")

    last_error: Optional[Exception] = None

    for encoding in encodings:
        try:
            return file_path.read_text(encoding=encoding)
        except (UnicodeDecodeError, OSError) as exc:
            last_error = exc

    if isinstance(last_error, OSError):
        raise last_error

    return file_path.read_text(encoding="utf-8", errors="replace")


def extract_text_file(file_path: Path) -> str:
    return read_text_file(file_path)


def extract_csv(file_path: Path) -> str:
    """
    Convert CSV rows to a simple pipe-separated representation.
    """
    encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    last_error: Optional[Exception] = None

    for encoding in encodings:
        try:
            with file_path.open(
                "r",
                encoding=encoding,
                newline="",
            ) as csv_file:
                reader = csv.reader(csv_file)
                rows = list(reader)

            return "\n".join(" | ".join(row) for row in rows)

        except (UnicodeDecodeError, OSError) as exc:
            last_error = exc

    if last_error:
        raise last_error

    return ""


def extract_pdf(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "PDF support requires 'pypdf'. " "Install it with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(file_path))
    pages: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"--- PAGE {index} ---\n{text.strip()}")

    return "\n\n".join(pages).strip()


def extract_docx(file_path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "Word support requires 'python-docx'. "
            "Install it with: pip install python-docx"
        ) from exc

    document = Document(str(file_path))
    sections: list[str] = []

    # Paragraphs
    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    if paragraphs:
        sections.append("\n".join(paragraphs))

    # Tables
    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []

        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))

        sections.append(f"--- TABLE {table_index} ---\n" + "\n".join(rows))

    return "\n\n".join(sections).strip()


def extract_xlsx(file_path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel support requires 'openpyxl'. "
            "Install it with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(
        filename=str(file_path),
        data_only=True,
        read_only=True,
    )

    sections: list[str] = []

    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []

            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]

                # Skip entirely blank rows.
                if any(value != "" for value in values):
                    rows.append(" | ".join(values))

            sections.append(
                f"--- SHEET: {sheet.title} ---\n"
                + ("\n".join(rows) if rows else "[Empty sheet]")
            )
    finally:
        workbook.close()

    return "\n\n".join(sections).strip()


def extract_pptx(file_path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "PowerPoint support requires 'python-pptx'. "
            "Install it with: pip install python-pptx"
        ) from exc

    presentation = Presentation(str(file_path))
    slides: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        texts: list[str] = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    texts.append(text)

        slides.append(
            f"--- SLIDE {slide_number} ---\n"
            + ("\n".join(texts) if texts else "[No extractable text]")
        )

    return "\n\n".join(slides).strip()


def extract_content(file_path: Path) -> str:
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".xlsx":
        return extract_xlsx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    if extension == ".csv":
        return extract_csv(file_path)

    if extension in TEXT_EXTENSIONS:
        return extract_text_file(file_path)

    raise ValueError(f"Unsupported file extension: {extension}")


# ============================================================
# Output path
# ============================================================


def get_valid_output_path() -> Path:
    while True:
        print("\nEnter the output file path:")
        print(r"Example: C:\Users\YourName\Desktop\extracted_content.txt")

        raw_path = input("> ")

        if not raw_path.strip():
            error("Output path cannot be empty.")
            continue

        path = clean_user_path(raw_path)

        # Give a useful error if the user accidentally enters a directory.
        if path.exists() and path.is_dir():
            error("The output path points to a directory, not a file.")
            continue

        if not path.suffix:
            warning("No file extension was provided.")
            print("Example: extracted_content.txt\n")
            continue

        try:
            path = path.resolve()
        except OSError:
            error("The output path could not be resolved.")
            continue

        parent = path.parent

        if not parent.exists():
            error(f"The parent directory does not exist:\n{parent}")
            print("Please enter an output path inside an existing folder.\n")
            continue

        if not parent.is_dir():
            error("The output parent path is not a directory.")
            continue

        if path.exists():
            print(f"\n[WARNING] This file already exists:\n{path}")
            print("\nDo you want to overwrite it?")
            print("1. Yes")
            print("2. No")

            choice = get_valid_choice("\nChoose an option: ", 1, 2)

            if choice == 2:
                print()
                continue

        return path


# ============================================================
# Output formatting
# ============================================================


def get_type_label(file_path: Path) -> str:
    extension = file_path.suffix.lower()
    name = SUPPORTED_EXTENSIONS.get(extension, "Unknown")
    return f"{name} ({extension})"


def build_file_block(
    source_folder: Path,
    file_number: int,
    file_path: Path,
    content: str,
) -> str:
    relative_path = file_path.relative_to(source_folder)
    type_label = get_type_label(file_path)

    return (
        f"\n{SEPARATOR}\n"
        f"FILE {file_number}\n"
        f"PATH: {relative_path}\n"
        f"TYPE: {type_label}\n"
        f"{SEPARATOR}\n\n"
        f"{content.rstrip()}\n\n"
    )


def write_output(
    source_folder: Path,
    selected_files: list[Path],
    output_file: Path,
) -> tuple[int, list[tuple[Path, str]]]:
    processed = 0
    skipped: list[tuple[Path, str]] = []

    with output_file.open("w", encoding="utf-8", newline="\n") as output:
        for file_number, file_path in enumerate(selected_files, start=1):
            try:
                content = extract_content(file_path)

                if not content.strip():
                    content = "[No extractable text/content found.]"

                output.write(
                    build_file_block(
                        source_folder,
                        file_number,
                        file_path,
                        content,
                    )
                )

                processed += 1

            except Exception as exc:
                reason = str(exc) or exc.__class__.__name__
                skipped.append((file_path, reason))

    return processed, skipped


# ============================================================
# Summary
# ============================================================


def display_summary(
    source_folder: Path,
    selected_files: list[Path],
    processed: int,
    skipped: list[tuple[Path, str]],
    output_file: Path,
) -> None:
    print("\n" + SEPARATOR)
    print(f"{'EXTRACTION COMPLETED':^70}")
    print(SEPARATOR)

    print(f"\nSource folder   : {source_folder}")
    print(f"Files selected  : {len(selected_files)}")
    print(f"Files processed : {processed}")
    print(f"Files skipped   : {len(skipped)}")
    print(f"Output file     : {output_file}")

    if skipped:
        print("\nSkipped files:")
        for file_path, reason in skipped:
            relative = file_path.relative_to(source_folder)
            print(f"- {relative} → {reason}")

    print()


# ============================================================
# Main application
# ============================================================


def main() -> None:
    print_header()

    # 1. Get source folder
    source_folder = get_valid_source_folder()

    print()
    info(f"Scanning folder recursively:\n{source_folder}")

    # 2. Scan recursively
    all_files = scan_files(source_folder)

    if not all_files:
        error("No files were found in the selected folder.")
        return

    info(f"Total files discovered: {len(all_files)}")

    # 3. Identify supported extensions
    extension_groups = get_extension_groups(all_files)

    if not extension_groups:
        error("No supported file types were found in the selected folder.")
        return

    # 4. Select type
    selected_by_type = select_file_type(extension_groups)

    if not selected_by_type:
        error("No files were selected.")
        return

    # 5. Display and select files
    selected_files = select_files(
        source_folder,
        selected_by_type,
    )

    if not selected_files:
        info("Operation cancelled.")
        return

    # 6. Ask for output path
    output_file = get_valid_output_path()

    # 7. Extract/write
    print("\n" + SEPARATOR)
    print("Starting extraction...")
    print(SEPARATOR)

    processed, skipped = write_output(
        source_folder,
        selected_files,
        output_file,
    )

    # 8. Final report
    display_summary(
        source_folder,
        selected_files,
        processed,
        skipped,
        output_file,
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\n[INFO] Operation cancelled by user.")
        sys.exit(0)
    except Exception as exc:
        print(f"\n[ERROR] Unexpected error: {exc}")
        sys.exit(1)
