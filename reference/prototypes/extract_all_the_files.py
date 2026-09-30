from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

# ============================================================
# SETTINGS
# ============================================================

source_folder = Path(r"C:\Java-script\session")
output_file = Path(r"C:\Java-script\new.txt")


# ============================================================
# FILE READERS
# ============================================================


def read_text_file(file_path):
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1")


def read_pdf(file_path):
    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n\n".join(pages)


def read_docx(file_path):
    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    return "\n".join(paragraphs)


def read_xlsx(file_path):
    workbook = load_workbook(file_path, data_only=True)

    result = []

    for sheet in workbook.worksheets:

        result.append(f"\n--- SHEET: {sheet.title} ---\n")

        for row in sheet.iter_rows(values_only=True):

            values = [str(value) if value is not None else "" for value in row]

            result.append(" | ".join(values))

    return "\n".join(result)


# ============================================================
# CONTENT EXTRACTION
# ============================================================


def extract_content(file_path):

    extension = file_path.suffix.lower()

    # Text-based files
    text_extensions = {
        ".py",
        ".txt",
        ".md",
        ".html",
        ".htm",
        ".css",
        ".js",
        ".jsx",
        ".ts",
        ".tsx",
        ".json",
        ".xml",
        ".csv",
        ".sql",
        ".java",
        ".c",
        ".cpp",
        ".h",
        ".hpp",
        ".php",
        ".rb",
        ".go",
        ".rs",
        ".sh",
        ".bat",
        ".yaml",
        ".yml",
        ".ini",
        ".log",
    }

    if extension in text_extensions:
        return read_text_file(file_path), "Text"

    # PDF
    if extension == ".pdf":
        return read_pdf(file_path), "PDF"

    # Word
    if extension == ".docx":
        return read_docx(file_path), "Word Document"

    # Excel
    if extension == ".xlsx":
        return read_xlsx(file_path), "Excel Spreadsheet"

    return None, "Unsupported"


# ============================================================
# COMBINE EVERYTHING
# ============================================================

files = sorted(
    file_path for file_path in source_folder.rglob("*") if file_path.is_file()
)

file_number = 0

with output_file.open("w", encoding="utf-8") as output:

    for file_path in files:

        content, file_type = extract_content(file_path)

        # Skip unsupported files
        if content is None:
            continue

        file_number += 1

        relative_path = file_path.relative_to(source_folder)

        output.write("\n")
        output.write("=" * 70)
        output.write("\n")

        output.write(f"{file_number}\n")
        output.write(f"FILE: {relative_path}\n")
        output.write(f"TYPE: {file_type}\n")

        output.write("=" * 70)
        output.write("\n\n")

        output.write(content)

        output.write("\n\n")


print(f"Done!")
print(f"Files included: {file_number}")
print(f"Output file: {output_file}")
